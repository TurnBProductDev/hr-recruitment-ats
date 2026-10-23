"""Recruitment manager's feedback (2026-10): Interview vs Interview Result
holds + a Hold Stage filter, Interviewer / Interview Score columns + an
Interviewer filter, "CV Score", mandatory remarks on every stage decision,
and the rejection stage + remarks on the Rejected list.

Run against sqlite so the live Azure DB is never touched:
    DB_ENGINE=sqlite python manage.py test candidates.test_recruiter_feedback
"""
import datetime
import importlib
import io

import openpyxl
from django.apps import apps
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from interviews.models import Interview, InterviewRequest

from . import services
from .models import Candidate, CandidateStatusHistory
from .permissions import HR_ADMIN

S = Candidate.Status
REPO = reverse('candidate_repository')


class _Base(TestCase):
    def setUp(self):
        User = get_user_model()
        self.user = User.objects.create_user('hr', 'hr@example.com', 'pw')
        self.user.groups.add(Group.objects.get_or_create(name=HR_ADMIN)[0])
        self.client.force_login(self.user)
        self.asha = User.objects.create_user('asha', 'asha@example.com', 'pw', first_name='Asha', last_name='Menon')
        self.ravi = User.objects.create_user('ravi', 'ravi@example.com', 'pw')
        self.candidate = self._candidate('Rose E G', 'rose@example.com')

    def _candidate(self, name, email, status=S.OPEN):
        c = Candidate.objects.create(full_name=name, email=email)
        services.record_creation(c)
        if status != S.OPEN:
            services.change_status(c, status, remarks='setup')
        return c

    def _interview(self, candidate, round_type=Interview.RoundType.ROUND1, status=Interview.Status.SCHEDULED,
                   interviewer=None, score=None):
        return Interview.objects.create(
            candidate=candidate, round_type=round_type, interviewer=interviewer or self.asha, status=status,
            score=score, scheduled_date=timezone.now() - datetime.timedelta(hours=1))

    def _post(self, url_name, candidate, **data):
        return self.client.post(reverse(url_name, args=[candidate.pk]), data)

    def _tab(self, tab, **params):
        return self.client.get(REPO, {'tab': tab, 'scoped': '1', **params})


class HoldKindTests(_Base):
    """Point 1 - a Round 1/Round 2 hold says whether the interview happened."""

    def test_hold_before_the_interview_is_an_interview_hold(self):
        services.change_status(self.candidate, S.ROUND1)
        self._interview(self.candidate, status=Interview.Status.CANCELLED)  # no-show
        self._post('candidate_screening_hold', self.candidate, reason='Candidate asked to postpone.')
        self.candidate.refresh_from_db()
        self.assertEqual(self.candidate.hold_kind, Candidate.HoldKind.INTERVIEW)
        self.assertEqual(self.candidate.status_label, 'Round 1 Interview Hold')

    def test_hold_after_the_interview_is_an_interview_result_hold(self):
        services.change_status(self.candidate, S.INTERVIEW)
        self._interview(self.candidate, Interview.RoundType.ROUND2, Interview.Status.COMPLETED)  # marked Done
        self._post('candidate_screening_hold', self.candidate, reason='Waiting on a second opinion.')
        self.candidate.refresh_from_db()
        self.assertEqual(self.candidate.status_label, 'Round 2 Interview Result Hold')

    def test_hold_picked_in_mark_result_is_an_interview_result_hold(self):
        services.change_status(self.candidate, S.ROUND1)
        interview = self._interview(self.candidate)
        self.client.post(reverse('interview_result', args=[interview.pk]), {
            'status': Interview.Status.COMPLETED, 'result': Interview.Result.HOLD, 'score': '7',
            'feedback': 'Good, but budget unclear.'})
        self.candidate.refresh_from_db()
        self.assertEqual(self.candidate.status_label, 'Round 1 Interview Result Hold')
        # The (mandatory) feedback is carried into the history remark too.
        self.assertIn('Good, but budget unclear.', self.candidate.history.latest('changed_at').remarks)

    def test_other_stages_keep_their_plain_hold_name(self):
        services.change_status(self.candidate, S.FINAL_SELECTION)
        self._post('candidate_screening_hold', self.candidate, reason='Offer on hold.')
        self.candidate.refresh_from_db()
        self.assertIsNone(self.candidate.hold_kind)
        self.assertEqual(self.candidate.status_label, 'Final Selection Hold')

    def test_activity_history_and_undo_keep_the_kind(self):
        services.change_status(self.candidate, S.ROUND1)
        self._interview(self.candidate, status=Interview.Status.COMPLETED)
        self._post('candidate_screening_hold', self.candidate, reason='Hold.')
        self._post('candidate_interview_stage', self.candidate, reason='Cleared after all.')
        response = self.client.get(reverse('candidate_timeline', args=[self.candidate.pk]))
        self.assertContains(response, 'Status: Round 1 → Round 1 Interview Result Hold')
        self.assertContains(response, 'Status: Round 1 Interview Result Hold → Round 2')
        self.client.post(reverse('candidate_revert', args=[self.candidate.pk]))  # undo Cleared
        self.candidate.refresh_from_db()
        self.assertEqual(self.candidate.status_label, 'Round 1 Interview Result Hold')

    def test_leaving_hold_clears_the_kind(self):
        services.change_status(self.candidate, S.ROUND1)
        services.change_status(self.candidate, S.SCREENING_HOLD)
        services.change_status(self.candidate, S.INTERVIEW)
        self.candidate.refresh_from_db()
        self.assertIsNone(self.candidate.hold_kind)

    def test_hold_stage_filter_on_the_hold_tab(self):
        before = self._candidate('Before', 'before@example.com', S.ROUND1)
        services.change_status(before, S.SCREENING_HOLD)
        after = self._candidate('After', 'after@example.com', S.ROUND1)
        self._interview(after, status=Interview.Status.COMPLETED)
        services.change_status(after, S.SCREENING_HOLD)
        final = self._candidate('Final', 'final@example.com', S.FINAL_SELECTION)
        services.change_status(final, S.SCREENING_HOLD)

        response = self._tab('screening_hold')
        self.assertContains(response, 'Hold Stage')
        self.assertContains(response, '<option value="ROUND1:RESULT"')
        self.assertEqual(len(response.context['candidates']), 3)
        listed = lambda value: [c.pk for c in self._tab('screening_hold', status=value).context['candidates']]
        self.assertEqual(listed('ROUND1:RESULT'), [after.pk])
        self.assertEqual(listed('ROUND1:INTERVIEW'), [before.pk])
        self.assertEqual(sorted(listed('ROUND1')), sorted([before.pk, after.pk]))
        self.assertEqual(listed('FINAL_SELECTION'), [final.pk])

    def test_backfill_names_existing_holds(self):
        result = self._candidate('Result', 'result@example.com', S.ROUND1)
        self._interview(result, status=Interview.Status.COMPLETED)
        services.change_status(result, S.SCREENING_HOLD)
        pending = self._candidate('Pending', 'pending@example.com', S.INTERVIEW)
        services.change_status(pending, S.SCREENING_HOLD)
        Candidate.objects.update(hold_kind=None)
        CandidateStatusHistory.objects.update(hold_kind=None)

        importlib.import_module('candidates.migrations.0024_hold_kind').backfill(apps, None)
        result.refresh_from_db()
        pending.refresh_from_db()
        self.assertEqual(result.hold_kind, Candidate.HoldKind.RESULT)
        self.assertEqual(pending.hold_kind, Candidate.HoldKind.INTERVIEW)


class InterviewColumnsTests(_Base):
    """Point 2 - Interviewer / Interview Score on Round 1, Round 2 and Hold."""

    def test_round1_tab_shows_interviewer_but_no_score_or_applied(self):
        services.change_status(self.candidate, S.ROUND1)
        self._interview(self.candidate, status=Interview.Status.COMPLETED, score='8.5')
        response = self._tab('round1')
        self.assertContains(response, '<th>Interviewer</th>')
        self.assertContains(response, 'Asha Menon')
        self.assertNotContains(response, '<th>Interview Score')
        self.assertNotContains(response, '<th>Applied</th>')

    def test_hold_tab_keeps_score_and_applied(self):
        services.change_status(self.candidate, S.ROUND1)
        self._interview(self.candidate, status=Interview.Status.COMPLETED, score='8.5')
        services.change_status(self.candidate, S.SCREENING_HOLD)
        response = self._tab('screening_hold')
        self.assertContains(response, '<th>Interview Score <span class="text-muted fw-normal">(/10)</span></th>')
        self.assertContains(response, '<strong>8.5</strong><span class="text-muted">/10</span>')
        self.assertContains(response, '<th>Applied</th>')

    def test_round2_tab_shows_previous_round_remarks(self):
        services.change_status(self.candidate, S.ROUND1)
        self._interview(self.candidate, status=Interview.Status.COMPLETED, score='9')
        Interview.objects.update(feedback='Strong on SQL.')
        services.change_status(self.candidate, S.INTERVIEW, remarks='Round 1 interview passed.')
        other = self._candidate('Other', 'other@example.com', S.ROUND1)
        services.change_status(other, S.INTERVIEW, remarks='Cleared without an interview.')
        response = self._tab('interview')
        self.assertContains(response, '<th>Previous Round Remarks</th>')
        self.assertNotContains(response, '<th>Applied</th>')
        rows = {c.pk: c.previous_round_remarks for c in response.context['candidates']}
        self.assertEqual(rows[self.candidate.pk], 'Round 1 interview passed. Strong on SQL.')
        self.assertEqual(rows[other.pk], 'Cleared without an interview.')

    def test_round_card_shows_interviewer_score_and_feedback(self):
        # Decided before Mark Result copied feedback into the remark - the
        # card still shows it, read from the interview itself.
        services.change_status(self.candidate, S.ROUND1)
        interview = self._interview(self.candidate, status=Interview.Status.COMPLETED, score='10')
        interview.result, interview.feedback = Interview.Result.PASS_, 'Okay'
        interview.save()
        services.change_status(self.candidate, S.INTERVIEW, remarks='Round 1 interview passed.')
        response = self.client.get(reverse('candidate_timeline', args=[self.candidate.pk]))
        self.assertContains(response, 'Interviewer: <strong>Asha Menon</strong>')
        self.assertContains(response, 'Interview Score: <strong>10</strong><span class="text-muted"> / 10</span>')
        self.assertContains(response, '<span class="text-muted">Feedback:</span> Okay')

    def test_round_card_does_not_repeat_feedback_already_in_the_remark(self):
        services.change_status(self.candidate, S.ROUND1)
        interview = self._interview(self.candidate)
        self.client.post(reverse('interview_result', args=[interview.pk]), {
            'status': Interview.Status.COMPLETED, 'result': Interview.Result.PASS_, 'score': '8',
            'feedback': 'Clear communicator.'})
        response = self.client.get(reverse('candidate_timeline', args=[self.candidate.pk]))
        self.assertContains(response, 'Round 1 interview passed. Clear communicator.')
        self.assertNotContains(response, '<span class="text-muted">Feedback:</span>')

    def test_round2_tab_shows_the_round2_interview_not_round1(self):
        services.change_status(self.candidate, S.INTERVIEW)
        self._interview(self.candidate, status=Interview.Status.COMPLETED, score='9')
        self._interview(self.candidate, Interview.RoundType.ROUND2, interviewer=self.ravi)
        row = self._tab('interview').context['candidates'][0]
        self.assertEqual((row.iv_interviewer, row.iv_score), ('ravi', None))

    def test_allocated_but_not_yet_scheduled_shows_the_allocated_interviewer(self):
        services.change_status(self.candidate, S.ROUND1)
        InterviewRequest.objects.create(candidate=self.candidate, interviewer=self.ravi)
        self.assertEqual(self._tab('round1').context['candidates'][0].iv_interviewer, 'ravi')

    def test_hold_tab_uses_the_interview_of_the_round_held_at(self):
        services.change_status(self.candidate, S.INTERVIEW)
        self._interview(self.candidate, status=Interview.Status.COMPLETED, score='6')
        self._interview(self.candidate, Interview.RoundType.ROUND2, Interview.Status.COMPLETED,
                        interviewer=self.ravi, score='7.5')
        services.change_status(self.candidate, S.SCREENING_HOLD)
        row = self._tab('screening_hold').context['candidates'][0]
        self.assertEqual((row.iv_interviewer, str(row.iv_score)), ('ravi', '7.5'))

    def test_interviewer_filter(self):
        services.change_status(self.candidate, S.ROUND1)
        self._interview(self.candidate)
        other = self._candidate('Other', 'other@example.com', S.ROUND1)
        self._interview(other, interviewer=self.ravi)
        response = self._tab('round1', interviewer=self.ravi.pk)
        self.assertEqual([c.pk for c in response.context['candidates']], [other.pk])
        self.assertIn((self.asha.pk, 'Asha Menon'), response.context['interviewer_options'])

    def test_columns_and_filter_not_on_other_tabs(self):
        response = self._tab('open')
        self.assertNotContains(response, '<th>Interviewer</th>')
        self.assertNotContains(response, 'All interviewers')


class CvScoreLabelTests(_Base):
    """Point 3 - the AI score is called CV Score."""

    def test_repository_and_export_say_cv_score(self):
        response = self._tab('open')
        self.assertContains(response, '<th>CV Score</th>')
        self.assertContains(response, 'All CV scores')
        export = self.client.get(reverse('candidate_repository_export'), {'tab': 'open', 'scoped': '1'})
        headers = [c.value for c in next(openpyxl.load_workbook(io.BytesIO(export.content)).active.iter_rows())]
        self.assertIn('CV Score', headers)
        self.assertNotIn('Score', headers)

    def test_candidate_page_says_cv_score(self):
        self.candidate.match_score = 72
        self.candidate.save(update_fields=['match_score'])
        response = self.client.get(reverse('candidate_timeline', args=[self.candidate.pk]))
        self.assertContains(response, 'CV Score')
        self.assertNotContains(response, 'Match Score</div>')


class MandatoryRemarksTests(_Base):
    """Point 4 - no stage decision without remarks."""

    def _assert_blocked(self, response, status=S.OPEN):
        self.candidate.refresh_from_db()
        self.assertEqual(self.candidate.status, status)
        return response

    def test_every_single_candidate_action_needs_remarks(self):
        for url_name in ('candidate_shortlist', 'candidate_screening_hold', 'candidate_reject',
                         'candidate_blacklist', 'candidate_move_to_future'):
            with self.subTest(url_name):
                response = self._post(url_name, self.candidate, reason='  ')
                self._assert_blocked(response)
                self.assertIn('Remarks are required', [str(m) for m in response.wsgi_request._messages][0])
        self._post('candidate_set_status', self.candidate, status=S.SHORTLISTED)
        self._assert_blocked(None)

    def test_with_remarks_it_goes_through(self):
        self._post('candidate_shortlist', self.candidate, reason='Strong CV.')
        self.candidate.refresh_from_db()
        self.assertEqual(self.candidate.status, S.SHORTLISTED)
        self.assertEqual(self.candidate.history.latest('changed_at').remarks, 'Strong CV.')

    def test_rejection_popup_needs_remarks(self):
        response = self.client.post(reverse('candidate_send_rejection', args=[self.candidate.pk]),
                                    {'skip_email': '1', 'to_email': ''}, HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(response.status_code, 400)
        self.assertIn('Remarks are required', response.json()['error'])
        self._assert_blocked(response)

    def test_bulk_actions_need_remarks(self):
        other = self._candidate('Other', 'other@example.com')
        for url_name in ('candidate_bulk_reject', 'candidate_bulk_shortlist', 'candidate_bulk_hold'):
            with self.subTest(url_name):
                self.client.post(reverse(url_name), {'ids': [self.candidate.pk, other.pk]})
                self.assertEqual(Candidate.objects.filter(status=S.OPEN).count(), 2)

    def test_candidate_page_gates_every_decision_button(self):
        services.change_status(self.candidate, S.SHORTLISTED)
        response = self.client.get(reverse('candidate_timeline', args=[self.candidate.pk]))
        content = response.content.decode()
        for url_name in ('candidate_round1', 'candidate_screening_hold'):
            url = reverse(url_name, args=[self.candidate.pk])
            tag = content[content.index(f'formaction="{url}"'):]
            self.assertIn('data-require-reason="1" disabled', tag[:tag.index('>')])
        # Logging an unanswered call is not a decision - stays clickable.
        self.assertNotIn('data-require-reason', content[content.index('value="UNABLE"'):][:200])


class RejectionListTests(_Base):
    """Point 6 - the Rejected tab says where and why."""

    def test_rejected_tab_shows_stage_and_remarks(self):
        services.change_status(self.candidate, S.ROUND1)
        self._post('candidate_reject', self.candidate, reason='Weak SQL fundamentals.')
        response = self._tab('rejected')
        self.assertContains(response, '<th>Rejected At</th><th>Remarks</th>')
        row = response.context['candidates'][0]
        self.assertEqual(row.rejection_stage, 'Round 1')
        self.assertContains(response, 'Weak SQL fundamentals.')

    def test_rejected_from_hold_names_the_hold(self):
        services.change_status(self.candidate, S.ROUND1)
        self._interview(self.candidate, status=Interview.Status.COMPLETED)
        services.change_status(self.candidate, S.SCREENING_HOLD)
        services.change_status(self.candidate, S.REJECTED, remarks='Position filled.')
        row = self._tab('rejected').context['candidates'][0]
        self.assertEqual(row.rejection_stage, 'Round 1 Interview Result Hold')

    def test_tele_screening_name_and_export(self):
        services.change_status(self.candidate, S.SHORTLISTED)
        services.change_status(self.candidate, S.REJECTED, remarks='Not interested.')
        export = self.client.get(reverse('candidate_repository_export'), {'tab': 'rejected', 'scoped': '1'})
        rows = list(openpyxl.load_workbook(io.BytesIO(export.content)).active.iter_rows(values_only=True))
        record = dict(zip(rows[0], rows[1]))
        self.assertEqual(record['Rejected At'], 'Tele Screening')
        self.assertEqual(record['Rejection Remarks'], 'Not interested.')

    def test_rejected_at_filter(self):
        services.change_status(self.candidate, S.ROUND1)
        services.change_status(self.candidate, S.REJECTED, remarks='x')
        tele = self._candidate('Tele', 'tele@example.com', S.SHORTLISTED)
        services.change_status(tele, S.REJECTED, remarks='x')
        held = self._candidate('Held', 'held@example.com', S.ROUND1)
        services.change_status(held, S.SCREENING_HOLD)
        services.change_status(held, S.REJECTED, remarks='x')

        response = self._tab('rejected')
        self.assertContains(response, 'Rejected At</label>')
        self.assertContains(response, '<option value="HOLD:ROUND1:INTERVIEW"')
        listed = lambda value: [c.pk for c in self._tab('rejected', rejected_at=value).context['candidates']]
        self.assertEqual(listed('ROUND1'), [self.candidate.pk])
        self.assertEqual(listed('SHORTLISTED'), [tele.pk])
        self.assertEqual(listed('HOLD:ROUND1:INTERVIEW'), [held.pk])
        self.assertEqual(listed('HOLD:ROUND1'), [held.pk])
        self.assertEqual(listed('HOLD:ROUND1:RESULT'), [])

    def test_no_applied_date_filter_on_rejected_hold_and_rounds(self):
        services.change_status(self.candidate, S.REJECTED, remarks='x')
        for tab in ('rejected', 'screening_hold', 'round1', 'interview'):
            with self.subTest(tab):
                self.assertNotContains(self._tab(tab), 'name="date_from"')
        self.assertContains(self._tab('open'), 'name="date_from"')
        # A date carried over from another tab doesn't silently filter it.
        response = self._tab('rejected', date_from='2099-01-01')
        self.assertEqual([c.pk for c in response.context['candidates']], [self.candidate.pk])

    def test_no_history_shows_a_dash(self):
        Candidate.objects.create(full_name='Legacy', email='legacy@example.com', status=S.REJECTED)
        response = self._tab('rejected')
        self.assertEqual(response.context['candidates'][0].rejection_stage, '')
